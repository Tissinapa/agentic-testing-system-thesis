"""
metrics/parse_schemathesis_results.py
Parse Schemathesis JUnit XML results into a clean Excel workbook.
Run from project root: python metrics/parse_schemathesis_results.py
"""

import re
import xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RESULTS_DIR = Path("results/schemathesis")
OUTPUT_PATH = Path("results/schemathesis_results.xlsx")


def thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)


def header_cell(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font      = Font(name="Calibri", bold=True, size=11)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border    = thin_border()


def data_cell(ws, row, col, value, center=False, bold=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name="Calibri", bold=bold, size=11)
    c.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center", wrap_text=True
    )
    c.border = thin_border()


def set_widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w


def clean_xml(content):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', content)


def parse_xml_file(filepath):
    content = filepath.read_text(encoding="utf-8", errors="replace")
    content = clean_xml(content)

    try:
        root = ET.fromstring(content)
    except ET.ParseError as e:
        print(f"Warning: XML parse error in {filepath.name}: {e}")
        tests    = int(re.search(r'tests="(\d+)"',    content).group(1)) if re.search(r'tests="(\d+)"',    content) else 0
        failures = int(re.search(r'failures="(\d+)"', content).group(1)) if re.search(r'failures="(\d+)"', content) else 0
        errors   = int(re.search(r'errors="(\d+)"',   content).group(1)) if re.search(r'errors="(\d+)"',   content) else 0
        return {
            "app": filepath.parent.name, "total": tests,
            "passed": max(0, tests - failures - errors),
            "failures": failures, "errors": errors,
            "cases": [], "parse_error": str(e),
        }

    total    = int(root.get("tests",    0))
    failures = int(root.get("failures", 0))
    errors   = int(root.get("errors",   0))
    passed   = max(0, total - failures - errors)

    cases = []
    for tc in root.iter("testcase"):
        failure = tc.find("failure")
        error   = tc.find("error")
        if failure is not None:
            status, message, detail = "FAIL", failure.get("message", ""), (failure.text or "").strip()[:500]
        elif error is not None:
            status, message, detail = "ERROR", error.get("message", ""), (error.text or "").strip()[:500]
        else:
            status, message, detail = "PASS", "", ""

        cases.append({
            "name": tc.get("name", ""), "classname": tc.get("classname", ""),
            "status": status, "message": message, "detail": detail, "time": tc.get("time", ""),
        })

    return {
        "app": filepath.parent.name, "total": total, "passed": passed,
        "failures": failures, "errors": errors, "cases": cases, "parse_error": None,
    }


def build_summary_sheet(wb, all_data):
    ws = wb.create_sheet("Summary", 0)
    ws.freeze_panes = "A2"
    headers = ["App", "Total Tests", "Passed", "Failures", "Errors", "Total Issues", "Pass Rate %"]
    for i, h in enumerate(headers):
        header_cell(ws, 1, i + 1, h)

    for i, d in enumerate(all_data):
        row       = i + 2
        total     = d["total"]
        issues    = d["failures"] + d["errors"]
        pass_rate = round(d["passed"] / total * 100, 1) if total > 0 else 0
        for j, v in enumerate([d["app"], total, d["passed"], d["failures"], d["errors"], issues, f"{pass_rate}%"]):
            data_cell(ws, row, j + 1, v, center=isinstance(v, (int, float)), bold=(j == 5))
        ws.row_dimensions[row].height = 16

    set_widths(ws, [16, 13, 10, 10, 10, 13, 12])


def build_issues_sheet(wb, all_data):
    ws = wb.create_sheet("All Issues")
    ws.freeze_panes = "A2"
    headers = ["App", "Test Name", "Status", "Message", "Detail"]
    for i, h in enumerate(headers):
        header_cell(ws, 1, i + 1, h)

    row = 2
    for d in all_data:
        for case in d["cases"]:
            if case["status"] in ("FAIL", "ERROR"):
                data_cell(ws, row, 1, d["app"],       bold=True)
                data_cell(ws, row, 2, case["name"])
                data_cell(ws, row, 3, case["status"], center=True, bold=True)
                data_cell(ws, row, 4, case["message"])
                data_cell(ws, row, 5, case["detail"])
                ws.row_dimensions[row].height = 50
                row += 1

    if row == 2:
        ws.cell(row=2, column=1, value="No issues found.")
    set_widths(ws, [16, 40, 8, 35, 55])


def build_detail_sheet(wb, data):
    app  = data["app"]
    name = app[:31]
    existing = [s.title for s in wb.worksheets]
    counter  = 1
    while name in existing:
        name = f"{app[:28]}_{counter}"
        counter += 1

    ws = wb.create_sheet(name)
    ws.freeze_panes = "A3"

    meta = (f"App: {app}  |  Total: {data['total']}  |  Passed: {data['passed']}  |  "
            f"Failures: {data['failures']}  |  Errors: {data['errors']}")
    ws.merge_cells(f"A1:{get_column_letter(5)}1")
    m = ws.cell(row=1, column=1, value=meta)
    m.font      = Font(name="Calibri", size=10, italic=True)
    m.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 16

    headers = ["Test Name", "Status", "Message", "Detail", "Time (s)"]
    for i, h in enumerate(headers):
        header_cell(ws, 2, i + 1, h)
    ws.row_dimensions[2].height = 20

    if not data["cases"]:
        ws.cell(row=3, column=1, value="No test case details available.")
        return

    for i, case in enumerate(data["cases"]):
        row  = i + 3
        bold = case["status"] in ("FAIL", "ERROR")
        data_cell(ws, row, 1, case["name"],    bold=bold)
        data_cell(ws, row, 2, case["status"],  center=True, bold=bold)
        data_cell(ws, row, 3, case["message"])
        data_cell(ws, row, 4, case["detail"])
        data_cell(ws, row, 5, case["time"],    center=True)
        ws.row_dimensions[row].height = 50

    set_widths(ws, [40, 8, 35, 55, 10])


def main():
    if not RESULTS_DIR.exists():
        print("No results/schemathesis directory found.")
        return

    xml_files = sorted(RESULTS_DIR.rglob("*.xml"))
    if not xml_files:
        print("No Schemathesis XML files found.")
        return

    all_data = []
    for f in xml_files:
        print(f"Parsing: {f.relative_to(RESULTS_DIR)}")
        data = parse_xml_file(f)
        all_data.append(data)
        print(f"  total: {data['total']}, issues: {data['failures'] + data['errors']}")

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    build_summary_sheet(wb, all_data)
    build_issues_sheet(wb, all_data)
    for d in all_data:
        build_detail_sheet(wb, d)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(f"\nSaved: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()